"""
Test all 5 generators with a tiny hard-coded structured dict.
No image, no OCR needed.
"""

import csv
from pathlib import Path
import pytest

from app.generators.convert import convert

# Minimal structured data that looks like a shopping list
SAMPLE = {
    "type": "list",
    "headers": ["qty", "unit", "item"],
    "data": [
        {"qty": "2",  "unit": "kg",  "item": "Rice",  "conf": 0.95},
        {"qty": "1",  "unit": "L",   "item": "Milk",  "conf": 0.60},  # low conf -> highlight
        {"qty": "6",  "unit": "",    "item": "Eggs",  "conf": 0.98},
    ],
    "confs": [0.95, 0.60, 0.98],
}

OUT = Path("data/test_output")


@pytest.fixture(autouse=True)
def cleanup():
    yield
    import shutil
    if OUT.exists():
        shutil.rmtree(OUT)


def test_csv():
    path = convert(SAMPLE, "csv", OUT)
    assert path.exists()
    rows = list(csv.reader(open(path, encoding="utf-8")))
    assert rows[0] == ["qty", "unit", "item"]
    assert rows[1] == ["2", "kg", "Rice"]
    assert rows[3] == ["6", "", "Eggs"]


def test_txt():
    path = convert(SAMPLE, "txt", OUT)
    assert path.exists()
    text = path.read_text(encoding="utf-8")
    assert "Rice" in text
    assert "Milk" in text


def test_excel():
    import openpyxl
    path = convert(SAMPLE, "excel", OUT)
    assert path.exists()
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws.cell(1, 1).value == "qty"    # header
    assert ws.cell(2, 3).value == "Rice"   # first data row, item column


def test_word():
    from docx import Document
    path = convert(SAMPLE, "word", OUT)
    assert path.exists()
    doc = Document(path)
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "SnapDocc" in full_text


def test_pdf():
    path = convert(SAMPLE, "pdf", OUT)
    assert path.exists()
    assert path.stat().st_size > 1000    # non-trivial file


def test_bad_format():
    with pytest.raises(ValueError):
        convert(SAMPLE, "pptx", OUT)
