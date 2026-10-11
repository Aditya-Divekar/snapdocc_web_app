"""Generate a Word (.docx) file from structured data."""

from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

LOW_CONF_COLOR = RGBColor(0xFF, 0xCC, 0x00)   # amber for low-confidence text
CONF_THRESHOLD = 0.75


def generate_word(structured, output_path):
    doc = Document()
    doc.add_heading("SnapDocc – Digitized Document", level=1)

    headers = structured.get("headers", [])
    data    = structured.get("data", [])

    # Add a table (works for both list and table doc types)
    n_cols = len(headers) if headers else 1
    table  = doc.add_table(rows=1, cols=n_cols)
    table.style = "Table Grid"

    # Header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        run = hdr_cells[i].paragraphs[0].runs[0]
        run.bold = True

    # Data rows
    for row in data:
        if structured["type"] == "list":
            cells     = [row.get("qty", ""), row.get("unit", ""), row.get("item", "")]
            conf_vals = [row.get("conf", 1.0)] * 3
        else:
            cells     = row["cells"]
            conf_vals = row["confs"]

        row_cells = table.add_row().cells
        for i, (value, conf) in enumerate(zip(cells, conf_vals)):
            para = row_cells[i].paragraphs[0]
            run  = para.add_run(str(value))
            if conf < CONF_THRESHOLD:
                run.font.color.rgb = LOW_CONF_COLOR   # amber = check me

    out = Path(output_path)
    doc.save(out)
    return out
