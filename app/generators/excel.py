"""
Generate an Excel (.xlsx) file from structured data.

Special feature: cells with confidence below CONF_THRESHOLD are highlighted
in yellow so the user knows to double-check them.
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from .sanitize import sanitize_cell

CONF_THRESHOLD = 0.75          # below this -> yellow highlight
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_headers(ws, headers):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=sanitize_cell(h))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def generate_excel(structured, output_path):
    """
    structured : dict returned by app.structuring.pipeline.structure()
    output_path: where to save the .xlsx file (string or Path)
    Returns    : Path object of the saved file
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SnapDocc"

    headers = structured.get("headers", [])
    data    = structured.get("data", [])
    confs   = structured.get("confs", [])

    _write_headers(ws, headers)

    for row_i, row in enumerate(data, 2):          # row 1 = headers
        # row is either a dict (list type) or a grid row dict (table type)
        if structured["type"] == "list":
            cells = [row.get("qty", ""), row.get("unit", ""), row.get("item", "")]
            conf_vals = [row.get("conf", 1.0)] * 3
        else:
            cells     = row["cells"]
            conf_vals = row["confs"]

        for col_i, (value, conf) in enumerate(zip(cells, conf_vals), 1):
            cell = ws.cell(row=row_i, column=col_i, value=sanitize_cell(value))
            if conf < CONF_THRESHOLD:
                cell.fill = YELLOW   # flag low-confidence cell

    # Auto-fit column widths (rough estimate)
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    out = Path(output_path)
    wb.save(out)
    return out
